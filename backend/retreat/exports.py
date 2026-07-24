"""
exports.py — professional Excel (openpyxl) and PDF (reportlab) builders.

pip install openpyxl reportlab

Every builder returns an in-memory BytesIO buffer, seeked to 0, ready to
stream straight back in an HttpResponse. Styling matches the app's palette
(navy / maroon / gold on cream) so exports look like part of the product,
not a raw data dump.
"""
import io
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.units import cm
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import (
    SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
)

NAVY   = "1C2541"
MAROON = "6E2C3A"
GOLD   = "D4A857"
STRIPE = "F4F0E6"

# ── Shared Excel helpers ──────────────────────────────────────────────────────

def _title_block(ws, title, subtitle=None):
    ws["A1"] = title
    ws["A1"].font = Font(size=16, bold=True, color=NAVY)
    ws["A2"] = subtitle or f"Generated {datetime.now():%d %b %Y, %H:%M}"
    ws["A2"].font = Font(size=9, italic=True, color="666666")


def _header_row(ws, row_idx, headers):
    fill = PatternFill(start_color=NAVY, end_color=NAVY, fill_type="solid")
    font = Font(color="FFFFFF", bold=True, size=11)
    thin = Side(style="thin", color="DDDDDD")
    for col, h in enumerate(headers, start=1):
        cell = ws.cell(row=row_idx, column=col, value=h)
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = Border(bottom=thin)


def _stripe_row(ws, row_idx, ncols):
    if row_idx % 2 == 0:
        fill = PatternFill(start_color=STRIPE, end_color=STRIPE, fill_type="solid")
        for c in range(1, ncols + 1):
            ws.cell(row=row_idx, column=c).fill = fill


def _autofit(ws, min_width=10, max_width=45):
    for col_cells in ws.columns:
        length = max((len(str(c.value)) for c in col_cells if c.value is not None), default=0)
        ws.column_dimensions[get_column_letter(col_cells[0].column)].width = min(
            max(length + 3, min_width), max_width
        )


def _finish(wb):
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


# ── Excel: participant registry ───────────────────────────────────────────────

def build_participants_excel(participants, title="Participant Registry"):
    wb = Workbook()
    ws = wb.active
    ws.title = "Participants"
    _title_block(ws, title, f"{len(participants)} participants · Generated {datetime.now():%d %b %Y, %H:%M}")

    headers = ["Name", "School", "Phone", "Category", "Sex", "Code"]
    header_row = 4
    _header_row(ws, header_row, headers)

    row = header_row + 1
    for p in participants:
        ws.cell(row=row, column=1, value=p.full_name)
        ws.cell(row=row, column=2, value=p.school)
        ws.cell(row=row, column=3, value=p.phone_number)
        ws.cell(row=row, column=4, value=p.category)
        ws.cell(row=row, column=5, value="Male" if p.sex == "M" else "Female")
        ws.cell(row=row, column=6, value=p.code)
        _stripe_row(ws, row, len(headers))
        row += 1

    ws.freeze_panes = f"A{header_row + 1}"
    _autofit(ws)
    return _finish(wb)


# ── Excel: single-session attendance sheet ────────────────────────────────────

def build_attendance_excel(session, records, title=None):
    wb = Workbook()
    ws = wb.active
    ws.title = "Attendance"
    program = session.retreat_day.program
    _title_block(
        ws,
        title or f"Attendance — {session.title}",
        f"{program.name} · Day {session.retreat_day.day_number} ({session.retreat_day.date}) · "
        f"Generated {datetime.now():%d %b %Y, %H:%M}",
    )

    headers = ["Name", "School", "Category", "Sex", "Present"]
    header_row = 4
    _header_row(ws, header_row, headers)

    row = header_row + 1
    present_fill = PatternFill(start_color="E7F5EC", end_color="E7F5EC", fill_type="solid")
    absent_fill = PatternFill(start_color="FBEAEC", end_color="FBEAEC", fill_type="solid")
    for r in records:
        ws.cell(row=row, column=1, value=r.participant.full_name)
        ws.cell(row=row, column=2, value=r.participant.school)
        ws.cell(row=row, column=3, value=r.participant.category)
        ws.cell(row=row, column=4, value="Male" if r.participant.sex == "M" else "Female")
        cell = ws.cell(row=row, column=5, value="Present" if r.present else "Absent")
        cell.fill = present_fill if r.present else absent_fill
        cell.font = Font(bold=True, color="1B7A43" if r.present else "A5313E")
        row += 1

    ws.freeze_panes = f"A{header_row + 1}"
    _autofit(ws)
    return _finish(wb)


# ── Excel: full program report (one sheet per day + summary) ─────────────────

def build_program_report_excel(program, days_data, total_registrations, unique_attendees):
    wb = Workbook()
    summary = wb.active
    summary.title = "Summary"
    _title_block(summary, f"{program.name} — Retreat Report", f"{program.start_date} → {program.end_date}")

    summary["A4"], summary["B4"] = "Total registrations", total_registrations
    summary["A5"], summary["B5"] = "Unique attendees", unique_attendees
    summary["A6"], summary["B6"] = "Days", len(days_data)
    for r in (4, 5, 6):
        summary.cell(row=r, column=1).font = Font(bold=True, color=NAVY)
    _autofit(summary)

    for day in days_data:
        ws = wb.create_sheet(title=f"Day {day['day_number']}"[:31])
        _title_block(ws, f"Day {day['day_number']} — {day['date']}", day.get("label", ""))

        ws["A4"], ws["B4"] = "Total present (unique)", day["total_present"]
        ws["A5"], ws["B5"] = "Registrations that day", day["registrations_today"]
        for r in (4, 5):
            ws.cell(row=r, column=1).font = Font(bold=True, color=NAVY)

        headers = ["Session", "Speaker", "Start", "Present", "Absent", "Total"]
        header_row = 7
        _header_row(ws, header_row, headers)
        row = header_row + 1
        for s in day["sessions"]:
            ws.cell(row=row, column=1, value=s["title"])
            ws.cell(row=row, column=2, value=s["speaker"])
            ws.cell(row=row, column=3, value=s["start_time"])
            ws.cell(row=row, column=4, value=s["present"])
            ws.cell(row=row, column=5, value=s["absent"])
            ws.cell(row=row, column=6, value=s["total"])
            _stripe_row(ws, row, len(headers))
            row += 1
        _autofit(ws)

    return _finish(wb)


# ── PDF helpers ────────────────────────────────────────────────────────────────

def _pdf_styles():
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(
        name="DLCFTitle", fontSize=20, textColor=colors.HexColor(f"#{NAVY}"),
        spaceAfter=4, fontName="Helvetica-Bold",
    ))
    styles.add(ParagraphStyle(
        name="DLCFSubtitle", fontSize=9, textColor=colors.HexColor("#666666"),
        spaceAfter=16,
    ))
    return styles


def _styled_table(data, col_widths=None):
    t = Table(data, colWidths=col_widths, repeatRows=1)
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(f"#{NAVY}")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 9),
        ("ALIGN", (0, 0), (-1, 0), "CENTER"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor(f"#{STRIPE}")]),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#DDDDDD")),
        ("TOPPADDING", (0, 0), (-1, -1), 6),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
        ("LEFTPADDING", (0, 0), (-1, -1), 8),
    ]))
    return t


def build_participants_pdf(participants, title="Participant Registry"):
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, topMargin=1.5 * cm, bottomMargin=1.5 * cm)
    styles = _pdf_styles()
    elements = [
        Paragraph(title, styles["DLCFTitle"]),
        Paragraph(f"{len(participants)} participants · Generated {datetime.now():%d %b %Y, %H:%M}", styles["DLCFSubtitle"]),
    ]
    data = [["Name", "School", "Phone", "Category", "Sex", "Code"]]
    for p in participants:
        data.append([p.full_name, p.school, p.phone_number, p.category, "M" if p.sex == "M" else "F", p.code])
    elements.append(_styled_table(data, col_widths=[4.2 * cm, 4.2 * cm, 2.8 * cm, 2.5 * cm, 1.3 * cm, 2.5 * cm]))
    doc.build(elements)
    buf.seek(0)
    return buf


def build_attendance_pdf(session, records):
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, topMargin=1.5 * cm, bottomMargin=1.5 * cm)
    styles = _pdf_styles()
    program = session.retreat_day.program
    present_count = sum(1 for r in records if r.present)

    elements = [
        Paragraph(f"Attendance — {session.title}", styles["DLCFTitle"]),
        Paragraph(
            f"{program.name} · Day {session.retreat_day.day_number} ({session.retreat_day.date}) · "
            f"{present_count} present of {len(records)} · Generated {datetime.now():%d %b %Y, %H:%M}",
            styles["DLCFSubtitle"],
        ),
    ]
    data = [["Name", "School", "Category", "Sex", "Status"]]
    for r in records:
        data.append([
            r.participant.full_name, r.participant.school, r.participant.category,
            "M" if r.participant.sex == "M" else "F", "Present" if r.present else "Absent",
        ])
    elements.append(_styled_table(data, col_widths=[4.5 * cm, 4.5 * cm, 2.8 * cm, 1.3 * cm, 2.6 * cm]))
    doc.build(elements)
    buf.seek(0)
    return buf


def build_program_report_pdf(program, days_data, total_registrations, unique_attendees):
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), topMargin=1.5 * cm, bottomMargin=1.5 * cm)
    styles = _pdf_styles()

    elements = [
        Paragraph(f"{program.name} — Retreat Report", styles["DLCFTitle"]),
        Paragraph(
            f"{program.start_date} → {program.end_date} · {total_registrations} registrations · "
            f"{unique_attendees} unique attendees · Generated {datetime.now():%d %b %Y, %H:%M}",
            styles["DLCFSubtitle"],
        ),
    ]

    for day in days_data:
        elements.append(Paragraph(
            f"Day {day['day_number']} — {day['date']}  ({day['total_present']} present, "
            f"{day['registrations_today']} registered that day)",
            ParagraphStyle(name="DayHeader", fontSize=12, textColor=colors.HexColor(f"#{MAROON}"),
                           spaceBefore=14, spaceAfter=6, fontName="Helvetica-Bold"),
        ))
        data = [["Session", "Speaker", "Start", "Present", "Absent", "Total"]]
        for s in day["sessions"]:
            data.append([s["title"], s["speaker"], s["start_time"], s["present"], s["absent"], s["total"]])
        elements.append(_styled_table(data, col_widths=[6 * cm, 5 * cm, 2.5 * cm, 2.5 * cm, 2.5 * cm, 2.5 * cm]))
        elements.append(Spacer(1, 6))

    doc.build(elements)
    buf.seek(0)
    return buf